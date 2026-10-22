import asyncio
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from vikram import documents


def fixture_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


class DocumentTests(unittest.TestCase):
    def test_generated_labels_do_not_make_empty_documents_readable(self):
        book = Workbook()
        book.active["A1"] = "  "
        xlsx = io.BytesIO()
        book.save(xlsx)
        doc = Document()
        doc.add_table(rows=2, cols=2)
        docx = io.BytesIO()
        doc.save(docx)
        for name, data in [
            ("empty.txt", b" \n\t\n"),
            ("empty.csv", b",,\n , , \n"),
            ("empty.xlsx", xlsx.getvalue()),
            ("empty.docx", docx.getvalue()),
        ]:
            with (
                self.subTest(name=name),
                self.assertRaisesRegex(ValueError, "readable"),
            ):
                documents.extract_document(name, data)

    def test_pdf_has_page_labels(self):
        result = documents.extract_document("hello.pdf", fixture_pdf())
        self.assertIn("Page 1", result["text"])
        self.assertIn("Hello PDF", result["text"])

    def test_word_preserves_paragraph_table_order(self):
        doc = Document()
        doc.add_paragraph("Before")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Inside"
        doc.add_paragraph("After")
        out = io.BytesIO()
        doc.save(out)
        text = documents.extract_document("hello.docx", out.getvalue())["text"]
        self.assertLess(text.index("Before"), text.index("Inside"))
        self.assertLess(text.index("Inside"), text.index("After"))

    def test_excel_labels_sheet_cells_and_unevaluated_formula(self):
        book = Workbook()
        book.active.title = "Budget"
        book.active.append(["Revenue", 12, "=B1*2"])
        out = io.BytesIO()
        book.save(out)
        text = documents.extract_document("budget.xlsx", out.getvalue())["text"]
        self.assertIn("Budget", text)
        self.assertIn("B1: 12", text)
        self.assertIn("unevaluated", text)

    def test_csv_quotes_and_utf16_text(self):
        text = documents.extract_document(
            "data.csv", b'name,description\n"a,b","line1\nline2"'
        )["text"]
        self.assertIn("a,b", text)
        self.assertIn("line1", text)
        self.assertIn(
            "hello",
            documents.extract_document("hello.txt", "hello".encode("utf-16"))["text"],
        )

    def test_text_clipping_binary_and_size(self):
        result = documents.extract_document("code.py", b"x" * 25000)
        self.assertLessEqual(len(result["text"]), 20000)
        self.assertIn("Partial", result["notice"])
        for name, data in [
            ("a.exe", b"bin"),
            ("a.txt", b"\x00\xff"),
            ("a.pdf", b"broken"),
            ("a.txt", b"x" * (10 * 1024 * 1024 + 1)),
        ]:
            with self.assertRaises(ValueError):
                documents.extract_document(name, data)

    def test_archive_expansion_limit_and_no_text_pdf(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("huge", b"x" * (51 * 1024 * 1024))
        with self.assertRaises(ValueError):
            documents.extract_document("a.docx", out.getvalue())
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        out = io.BytesIO()
        writer.write(out)
        with self.assertRaisesRegex(ValueError, "OCR"):
            documents.extract_document("scan.pdf", out.getvalue())


class WorkerTests(unittest.IsolatedAsyncioTestCase):
    async def test_worker_extracts_and_timeout_terminates(self):
        result = await documents.extract_in_worker("hello.txt", b"hello")
        self.assertIn("hello", result["text"])
        with patch("vikram.documents.EXTRACT_TIMEOUT", 0.001):
            with self.assertRaisesRegex(ValueError, "too long"):
                await documents.extract_in_worker("hello.txt", b"hello")
