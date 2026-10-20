"""Local bounded document extraction and disposable parser process ownership."""

import asyncio
import base64
import csv
import io
import json
import shutil
import sys
import zipfile
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 20_000
EXTRACT_TIMEOUT = 15
TEXT_EXTENSIONS = set(
    "txt md markdown rst log json jsonl yaml yml toml xml html htm css scss js jsx ts tsx py pyi rs go java c h cpp hpp cs swift kt kts m mm sh bash zsh fish sql rb php r jl lua pl ini cfg conf properties csv tsv".split()
)
MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def ensure_space(path: Path, required: int = 0) -> None:
    if shutil.disk_usage(path).free - required < 20 * 1024**3:
        raise ValueError(
            "This action needs more free space to keep the 20 GiB disk reserve."
        )


def decode_text(data: bytes) -> str:
    try:
        text = data.decode(
            "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        )
        if any(ord(c) < 32 and c not in "\n\r\t" for c in text):
            raise ValueError()
        return text
    except (UnicodeError, ValueError):
        raise ValueError(
            "Choose UTF-8 or BOM-marked UTF-16 text; binary files are unsupported."
        ) from None


class TextBuffer:
    def __init__(self):
        self.text = ""
        self.notices = []
        self.has_content = False

    def add(self, value: str, *, content: str | None = None) -> bool:
        remaining = MAX_TEXT - len(self.text)
        self.has_content |= bool(
            (value if content is None else content)[:remaining].strip()
        )
        self.text += value[:remaining]
        if len(value) > remaining:
            self.notices.append("Partial: extracted text reached 20,000 characters.")
            return False
        return True


def validate_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
            raise ValueError("This Office file exceeds the archive expansion limit.")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("Encrypted documents are unsupported.")


def extract_document(name: str, data: bytes) -> dict:
    ext = Path(name).suffix.lower().lstrip(".")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Attachments must be nonempty and no larger than 10 MiB.")
    if ext not in TEXT_EXTENSIONS | set(MEDIA):
        raise ValueError("Choose PDF, DOCX, XLSX, CSV, or a supported text/code file.")
    result = TextBuffer()
    try:
        if ext == "pdf":
            from pypdf import PdfReader, overwrite_configuration

            overwrite_configuration(zlib_maximum_output_length=50 * 1024 * 1024)
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are unsupported.")
            if len(reader.pages) > 100:
                result.notices.append(
                    "Partial: only the first 100 PDF pages were inspected."
                )
            empty = []
            found = False
            for index, page in enumerate(reader.pages[:100], 1):
                text = page.extract_text() or ""
                if not text.strip():
                    empty.append(str(index))
                    continue
                found = True
                if not result.add(f"Page {index}\n{text}\n\n"):
                    break
            if not found:
                raise ValueError(
                    "This PDF has no extractable text. OCR is not supported."
                )
            if empty:
                result.notices.append(
                    "Partial: no extractable text on pages "
                    + ", ".join(empty)
                    + "; OCR is not supported."
                )
        elif ext == "docx":
            from docx import Document
            from docx.table import Table

            validate_archive(data)
            for item in Document(io.BytesIO(data)).iter_inner_content():
                text = (
                    "\n".join(
                        " | ".join(c.text for c in row.cells) for row in item.rows
                    )
                    if isinstance(item, Table)
                    else item.text
                )
                content = (
                    "\n".join(c.text for row in item.rows for c in row.cells)
                    if isinstance(item, Table)
                    else text
                )
                if not result.add(text + "\n", content=content):
                    break
            result.notices.append(
                "Main-body text and tables; drawings, headers, footers, and embedded objects are not included."
            )
        elif ext == "xlsx":
            from openpyxl import load_workbook

            validate_archive(data)
            formulas = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            cached = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            try:
                if len(formulas.worksheets) > 10:
                    result.notices.append(
                        "Partial: only the first 10 sheets were inspected."
                    )
                for sheet in formulas.worksheets[:10]:
                    result.add(f"Sheet: {sheet.title}\n", content="")
                    if (sheet.max_row or 0) > 1000 or (sheet.max_column or 0) > 100:
                        result.notices.append(
                            "Partial: up to 1,000 rows and 100 columns per sheet."
                        )
                    values = cached[sheet.title].iter_rows(max_row=1000, max_col=100)
                    for row, value_row in zip(
                        sheet.iter_rows(max_row=1000, max_col=100), values
                    ):
                        cells = []
                        for cell, value_cell in zip(row, value_row):
                            if cell.value is None or not str(cell.value).strip():
                                continue
                            value = cell.value
                            if cell.data_type == "f":
                                value = (
                                    value_cell.value
                                    if value_cell.value is not None
                                    else f"[unevaluated formula: {value}]"
                                )
                            cells.append(f"{cell.coordinate}: {value}")
                        if cells and not result.add(" | ".join(cells) + "\n"):
                            break
                    if len(result.text) >= MAX_TEXT:
                        break
                result.notices.append(
                    "Formula values use saved caches where available; formulas are not recalculated."
                )
            finally:
                formulas.close()
                cached.close()
        else:
            text = decode_text(data)
            if ext in {"csv", "tsv"}:
                csv.field_size_limit(MAX_BYTES)
                rows = csv.reader(
                    io.StringIO(text), delimiter="\t" if ext == "tsv" else ","
                )
                for index, row in enumerate(rows, 1):
                    if index > 1000:
                        result.notices.append(
                            "Partial: only the first 1,000 rows were included."
                        )
                        break
                    if len(row) > 100:
                        result.notices.append(
                            "Partial: only the first 100 columns were included."
                        )
                    if not result.add(
                        f"Row {index}: "
                        + json.dumps(row[:100], ensure_ascii=False)
                        + "\n",
                        content="\n".join(row[:100]),
                    ):
                        break
            else:
                for index, line in enumerate(text.splitlines(), 1):
                    if not result.add(f"{index}: {line}\n", content=line):
                        break
    except ValueError:
        raise
    except Exception:
        raise ValueError("This document is damaged or cannot be read.") from None
    if not result.has_content:
        raise ValueError("This document has no readable content.")
    return {
        "kind": "document",
        "media_type": MEDIA.get(
            ext, "text/csv" if ext in {"csv", "tsv"} else "text/plain"
        ),
        "width": 0,
        "height": 0,
        "text": result.text,
        "notice": " ".join(dict.fromkeys(result.notices)) or None,
        "extraction_version": 1,
    }


async def extract_in_worker(name: str, data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise ValueError("Attachments must be no larger than 10 MiB.")
    command = (
        [sys.executable, "--extract-document"]
        if getattr(sys, "frozen", False)
        else [sys.executable, "-m", "vikram.document_worker"]
    )
    process = await asyncio.create_subprocess_exec(
        *command,
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        async with asyncio.timeout(EXTRACT_TIMEOUT):
            out, _ = await process.communicate(
                json.dumps(
                    {"name": name, "data": base64.b64encode(data).decode()}
                ).encode()
            )
        if process.returncode != 0 or len(out) > 512 * 1024:
            raise ValueError(
                "This document could not be extracted within its resource limits."
            )
        value = json.loads(out)
        if "error" in value:
            raise ValueError(value["error"])
        return value
    except TimeoutError:
        raise ValueError(
            "Document extraction took too long (15-second limit)."
        ) from None
    finally:
        if process.returncode is None:
            process.kill()
        await process.wait()
